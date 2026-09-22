from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT = 'tmp/הצהרת בריאות לנהג.docx'

def set_rtl(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    ppr.append(bidi)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), fill)
    tc_pr.append(shading)

def set_cell_border(cell, color='D9D9D9'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), '6')
        tag.set(qn('w:color'), color)
        borders.append(tag)
    tc_pr.append(borders)

def add_text(paragraph, text, bold=False, size=11, color=None):
    set_rtl(paragraph)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run

def remove_paragraph_borders(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn('w:pBdr'))
    if borders is not None:
        ppr.remove(borders)

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
style = doc.styles['Normal']
style.font.name = 'Arial'
style._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
style._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
style.font.size = Pt(11)
title_style = doc.styles['Title']
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = title_style._element.get_or_add_pPr()
title_borders = title_ppr.find(qn('w:pBdr'))
if title_borders is not None:
    title_ppr.remove(title_borders)

title = doc.add_paragraph(style='Title')
remove_paragraph_borders(title)
add_text(title, 'הצהרת בריאות לנהג', bold=True, size=19, color='000000')
subtitle = doc.add_paragraph()
add_text(subtitle, 'טופס הצהרה אישית לפני ביצוע עבודת נהיגה', size=11, color='404040')
intro = doc.add_paragraph()
add_text(intro, 'הטופס מיועד לנהג ולחברה המעסיקה. יש להשיב בכנות ובמלואן על השאלות שלהלן. הטופס אינו מהווה אבחון רפואי ואינו מחליף ייעוץ רפואי.', size=10)
doc.add_paragraph()

heading = doc.add_paragraph()
add_text(heading, 'פרטי הנהג והחברה', bold=True, size=13)
details = doc.add_table(rows=4, cols=2)
details.style = 'Table Grid'
for row, (label, value) in zip(details.rows, [
    ('שם החברה', '{{company_name}}'), ('שם הנהג', '{{driver_full_name}}'),
    ('טלפון', '{{driver_phone}}'), ('תעודת זהות', '{{driver_national_id}}'),
]):
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
    set_cell_shading(row.cells[0], 'F2F4F7')
    add_text(row.cells[0].paragraphs[0], label, bold=True, size=10)
    add_text(row.cells[1].paragraphs[0], value, size=10)

doc.add_paragraph()
questions_heading = doc.add_paragraph()
add_text(questions_heading, 'הצהרת הנהג', bold=True, size=13)
lead = doc.add_paragraph()
add_text(lead, 'אני מצהיר או מצהירה כי למיטב ידיעתי:', size=11)
for statement in [
    'אני כשיר או כשירה לבצע עבודת נהיגה באופן בטוח ואחראי.',
    'אין לי מגבלה רפואית ידועה שעלולה לפגוע ביכולת הנהיגה שלי, או שדיווחתי עליה לחברה כנדרש.',
    'אני נוטל או נוטלת תרופות רק בהתאם להנחיות, ואדווח אם לתרופה עלולה להיות השפעה על ערנות או נהיגה.',
    'אדווח לחברה בהקדם על שינוי במצבי הבריאותי שעלול להשפיע על בטיחות הנהיגה.',
    'הפרטים שמסרתי בטופס זה נכונים ומלאים למיטב ידיעתי.',
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_text(p, '☐  ' + statement, size=11)
notice = doc.add_paragraph()
notice.paragraph_format.space_before = Pt(8)
add_text(notice, 'אם יש ספק לגבי כשירות לנהיגה, יש להימנע מנהיגה ולפנות לגורם רפואי מוסמך או לממונה בחברה.', bold=True, size=10)

doc.add_paragraph()
signature_heading = doc.add_paragraph()
add_text(signature_heading, 'אישור וחתימה', bold=True, size=13)
signature = doc.add_table(rows=2, cols=2)
signature.style = 'Table Grid'
for row, (label, value) in zip(signature.rows, [('תאריך', '{{date}}'), ('חתימת הנהג', '{{signature}}')]):
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
    set_cell_shading(row.cells[0], 'F2F4F7')
    add_text(row.cells[0].paragraphs[0], label, bold=True, size=10)
    add_text(row.cells[1].paragraphs[0], value, size=10)
footer = section.footer.paragraphs[0]
add_text(footer, 'FleetOS  |  טיוטת תבנית לבדיקת התהליך  |  נדרש אישור החברה לפני שימוש תפעולי', size=8, color='666666')
doc.save(OUTPUT)
print(OUTPUT)
